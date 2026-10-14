import pickle
import streamlit as st
import time
from pptx import Presentation
import datetime
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
import io
import psycopg2

# Connects directly to Supabase using the credentials saved in Streamlit Secrets
def get_db_connection():
    db_url = st.secrets["postgres"]["url"]
    return psycopg2.connect(db_url)

def init_db():
    conn = get_db_connection()
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS presentations 
                 (id TEXT PRIMARY KEY, title TEXT, pptx_data BYTEA)''')
    c.execute('''CREATE TABLE IF NOT EXISTS questions 
                 (id SERIAL PRIMARY KEY, pres_id TEXT, 
                  question TEXT, op1 TEXT, op2 TEXT, op3 TEXT, op4 TEXT, correct INTEGER)''')
    conn.commit()
    c.close()
    conn.close()

init_db()

st.set_page_config(page_title="Multi-Course Training Portal", layout="centered", page_icon="🎓")

query_params = st.query_params
presentation_id = query_params.get("id", None)
if isinstance(presentation_id, list):
    presentation_id = presentation_id[0]
mode = query_params.get("mode", "student" if presentation_id else "admin")

# --- ADMIN DASHBOARD ---
if mode == "admin":
    st.title("🛡️ Training Portal - Admin Dashboard")
    st.write("Upload a PowerPoint presentation or slide images, and construct an associated 5-question multi-choice quiz.")
    
    admin_token = st.text_input("Enter Admin Password", type="password")
    if admin_token == "admin123":
        st.success("Access Verification Successful")
        
        with st.form("create_course"):
            course_id = st.text_input("Unique Course ID (e.g., safety-2026)", help="This forms part of the unique sharing URL link")
            course_title = st.text_input("Course Title")
            uploaded_files = st.file_uploader("Upload PowerPoint Presentation (.pptx) or Slide Images (.png, .jpg)", type=["pptx", "png", "jpg", "jpeg"], accept_multiple_files=True)
            
            st.write("---")
            st.subheader("📋 Quiz Setup (Configure 5 Custom Questions)")
            
            questions_data = []
            for i in range(1, 6):
                st.markdown(f"### Question {i}")
                q = st.text_input("Question text", key=f"q_{i}")
                o1 = st.text_input("Option A", key=f"o1_{i}")
                o2 = st.text_input("Option B", key=f"o2_{i}")
                o3 = st.text_input("Option C", key=f"o3_{i}")
                o4 = st.text_input("Option D", key=f"o4_{i}")
                correct = st.selectbox("Correct Option", ["A", "B", "C", "D"], key=f"c_{i}")
                correct_idx = ["A", "B", "C", "D"].index(correct) + 1
                questions_data.append((q, o1, o2, o3, o4, correct_idx))
            
            submit = st.form_submit_button("Publish Course Module")
            
            if submit:
                if not course_id or not course_title or not uploaded_files:
                    st.error("All visual configurations, text labels, and at least one valid file (.pptx, .png, .jpg) must be supplied.")
                else:
                    conn = get_db_connection()
                    c = conn.cursor()
                    try:
                        if len(uploaded_files) == 1 and uploaded_files[0].name.lower().endswith(".pptx"):
                            file_data = uploaded_files[0].read()
                        else:
                            sorted_files = sorted(uploaded_files, key=lambda x: x.name)
                            image_bytes_list = [f.read() for f in sorted_files]
                            file_data = pickle.dumps(image_bytes_list)
                        
                        c.execute("INSERT INTO presentations (id, title, pptx_data) VALUES (%s, %s, %s)",
                                  (course_id, course_title, psycopg2.Binary(file_data)))

                        for q_data in questions_data:
                            c.execute("INSERT INTO questions (pres_id, question, op1, op2, op3, op4, correct) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                      (course_id, *q_data))
                        conn.commit()
                        st.success("🎉 Training Module Successfully Stored Permanently in Supabase!")
                        
                        share_url = f"https://training-app-cpd.streamlit.app/?id={course_id}"
                        st.info(f"**Shareable Presentation URL for Students:** `{share_url}`")
                    except psycopg2.IntegrityError:
                        conn.rollback()
                        st.error("That Course ID already exists. Please choose a different unique identifier.")
                    finally:
                        c.close()
                        conn.close()
    elif admin_token:
        st.error("Invalid Administrative Credentials Provided")

# --- STUDENT TRAINING PORTAL ---
else:
    if not presentation_id:
        st.title("🎓 Institutional Corporate Learning Portal")
        st.warning("Please utilize the official direct resource link configured by your training program coordinator.")
    else:
        conn = get_db_connection()
        c = conn.cursor()
        c.execute("SELECT title, pptx_data FROM presentations WHERE id = %s", (presentation_id,))
        course = c.fetchone()
        
        if not course:
            st.error("Requested learning module could not be verified. Confirm your target hyperlink identifier.")
            c.close()
            conn.close()
            st.stop()
            
        course_title, pptx_bytes = course
        
        if isinstance(pptx_bytes, memoryview):
            pptx_bytes = bytes(pptx_bytes)

        c.execute("SELECT question, op1, op2, op3, op4, correct FROM questions WHERE pres_id = %s", (str(presentation_id),))
        quiz_questions = c.fetchall()
        c.close()
        conn.close()

        is_image_slides = False
        try:
            slides = pickle.loads(pptx_bytes)
            if isinstance(slides, list):
                is_image_slides = True
        except Exception:
            slides = pptx_bytes

        st.title(f"📖 Active Module: {course_title}")
        
        if "slide_index" not in st.session_state:
            st.session_state.slide_index = 0
        if "timer_start" not in st.session_state:
            st.session_state.timer_start = time.time()
        if "quiz_started" not in st.session_state:
            st.session_state.quiz_started = False
        if "passed" not in st.session_state:
            st.session_state.passed = False

        prs = None
        if is_image_slides:
            total_slides = len(slides)
        else:
            try:
                prs = Presentation(io.BytesIO(pptx_bytes))
                total_slides = len(prs.slides)
            except Exception:
                total_slides = 1

        # Phase 1: Slide Navigation
        if not st.session_state.quiz_started:
            current_slide = st.session_state.slide_index
            st.subheader(f"Presentation View: Slide {current_slide + 1} of {total_slides}")
            
            if is_image_slides:
                idx = min(max(0, current_slide), len(slides) - 1)
                st.image(io.BytesIO(slides[idx]), use_container_width=True)
            elif prs:
                slide = prs.slides[current_slide]
                text_runs = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                if text_runs:
                    st.info("\n\n".join(text_runs))
                else:
                    st.warning("Slide contains visual elements. Proceed using the timer controls below.")

            elapsed = time.time() - st.session_state.timer_start
            time_remaining = max(0, 3 - int(elapsed))
            
            if time_remaining > 0:
                st.button(f"⏱️ Next Slide Locked ({time_remaining}s remaining)", disabled=True)
                time.sleep(1)
                st.rerun()
            else:
                if current_slide < total_slides - 1:
                    if st.button("➡️ Next Slide"):
                        st.session_state.slide_index += 1
                        st.session_state.timer_start = time.time()
                        st.rerun()
                else:
                    if st.button("📝 Initiate Certification Assessment"):
                        st.session_state.quiz_started = True
                        st.rerun()

        # Phase 2: Assessment Phase
        elif st.session_state.quiz_started and not st.session_state.passed:
            st.subheader("📋 Assessment Phase")
            st.write("Achieve a perfect 100% score (5/5 answers correct) to finalize certification. Unlimited retries allowed.")
            
            with st.form("quiz_form"):
                user_answers = []
                for idx, q in enumerate(quiz_questions):
                    st.write(f"**Q{idx+1}: {q[0]}**")
                    ans = st.radio("Select choice:", [q[1], q[2], q[3], q[4]], key=f"student_q_{idx}")
                    ans_idx = [q[1], q[2], q[3], q[4]].index(ans) + 1
                    user_answers.append(ans_idx)
                
                submit_quiz = st.form_submit_button("Submit Evaluation")
                
                if submit_quiz:
                    correct_count = sum(1 for u, q in zip(user_answers, quiz_questions) if u == q[5])
                    if correct_count == len(quiz_questions) and len(quiz_questions) > 0:
                        st.session_state.passed = True
                        st.success("🎉 Verification Criteria Satisfied! Proceed to enter credentials.")
                        st.rerun()
                    else:
                        st.error(f"❌ Standard Not Met ({correct_count}/{len(quiz_questions)} correct). Review material and restart evaluation.")

        # Phase 3: Completion & Certificate Issuance
        elif st.session_state.passed:
            st.balloons()
            st.success("🎓 Training Requirements Successfully Satisfied!")
            
            student_name = st.text_input("Enter Legal Full Name for Document Issuance:", key="cert_name_input")
            
            if student_name.strip():
                buffer = io.BytesIO()
                doc = SimpleDocTemplate(
                    buffer, 
                    pagesize=landscape(letter),
                    rightMargin=40, leftMargin=40, topMargin=50, bottomMargin=40, 
                    title="Training Certificate"
                )
                styles = getSampleStyleSheet()
                
                title_style = ParagraphStyle('CertTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=32, leading=38, alignment=TA_CENTER)
                body_style = ParagraphStyle('CertBody', parent=styles['Normal'], fontName='Helvetica', fontSize=18, leading=24, alignment=TA_CENTER)
                name_style = ParagraphStyle('CertName', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=26, leading=32, alignment=TA_CENTER, textColor='navy')
                now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                
                story = [
                    Paragraph("CERTIFICATE OF COMPLETION", title_style),
                    Spacer(1, 30),
                    Paragraph("This record serves to verify that", body_style),
                    Spacer(1, 15),
                    Paragraph(student_name.strip().upper(), name_style),
                    Spacer(1, 15),
                    Paragraph("has fully complied with and finished the core instructional guidelines for", body_style),
                    Spacer(1, 10),
                    Paragraph(f"<b>{course_title}</b>", body_style),
                    Spacer(1, 30),
                    Paragraph(f"Timestamp of Issue: {now}", body_style),
                    Spacer(1, 20),
                    Paragraph("This is an officially issued digital completion certificate.", ParagraphStyle('Footer', parent=body_style, fontSize=9, textColor='gray'))
                ]
                
                doc.build(story)
                buffer.seek(0)
                
                st.download_button(
                    label="📥 Generate and Download Official PDF Certificate",
                    data=buffer.getvalue(),
                    file_name=f"Certificate_{student_name.strip().replace(' ', '_')}.pdf",
                    mime="application/pdf"
                )
